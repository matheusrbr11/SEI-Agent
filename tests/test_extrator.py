from extratores.extrator import _eh_pdf, extrair_conteudo


def test_texto_simples(tmp_path):
    arq = tmp_path / "nota.txt"
    arq.write_text("olá mundo", encoding="utf-8")
    assert extrair_conteudo(arq) == "olá mundo"


def test_extensao_desconhecida_retorna_vazio(tmp_path):
    arq = tmp_path / "binario.bin"
    arq.write_bytes(b"\x00\x01\x02")
    assert extrair_conteudo(arq) == ""


def test_arquivo_inexistente_retorna_vazio(tmp_path):
    assert extrair_conteudo(tmp_path / "nao_existe.pdf") == ""


def test_eh_pdf_por_assinatura(tmp_path):
    falso = tmp_path / "x.txt"
    falso.write_bytes(b"%PDF-1.4 conteudo")
    assert _eh_pdf(falso)
    assert not _eh_pdf(tmp_path / "inexistente")


def test_pdf_corrompido_nao_levanta_excecao(tmp_path):
    arq = tmp_path / "ruim.pdf"
    arq.write_bytes(b"%PDF-1.4 lixo")
    assert extrair_conteudo(arq) == ""


def test_docx(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Primeiro parágrafo")
    doc.add_paragraph("   ")
    doc.add_paragraph("Segundo parágrafo")
    arq = tmp_path / "a.docx"
    doc.save(str(arq))
    assert extrair_conteudo(arq) == "Primeiro parágrafo\nSegundo parágrafo"


def test_excel_limita_linhas(tmp_path):
    import pandas as pd
    from extratores.extrator import extrair_excel
    arq = tmp_path / "p.xlsx"
    pd.DataFrame({"a": range(150)}).to_excel(arq, index=False, sheet_name="Dados")
    texto = extrair_excel(arq, max_linhas=10)
    assert "### Aba: Dados (150 linhas)" in texto
    assert "9" in texto and "149" not in texto
